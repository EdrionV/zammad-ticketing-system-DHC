import pandas as pd
import openpyxl
import os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image
import subprocess

# 1. Execute PostgreSQL copy command via Docker to export raw ticket data to CSV
export_cmd = (
    "sudo docker exec -i zammad-docker-compose-zammad-postgresql-1 psql -U zammad -d zammad_production "
    "-c \"\\copy (SELECT ROW_NUMBER() OVER (ORDER BY t.id ASC) AS id, t.number AS ticket_number, t.title, "
    "u.firstname || ' ' || u.lastname AS customer_name, t.state_id, t.priority_id, "
    "TO_CHAR(t.created_at, 'YYYY-MM-DD') AS created_date, TO_CHAR(t.created_at, 'HH24:MI:SS') AS created_time "
    "FROM tickets t LEFT JOIN users u ON t.customer_id = u.id ORDER BY t.id ASC) "
    "TO STDOUT WITH CSV HEADER\" > ~/tickets_export.csv"
)
subprocess.run(export_cmd, shell=True, check=True)

# 2. Read CSV and initialise workbook
csv_path = '/home/ed/tickets_export.csv'
excel_path = '/home/ed/tickets_export.xlsx'

df = pd.read_csv(csv_path, dtype={'ticket_number': str})
df.to_excel(excel_path, index=False)

wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# 3. Define DHC Subtle Theme Palette
header_fill = PatternFill(start_color='7A1F5C', end_color='7A1F5C', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
zebra_fill = PatternFill(start_color='FDF2F7', end_color='FDF2F7', fill_type='solid')
white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
border_side = Side(border_style='thin', color='E0C2D8')
cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
cell_font = Font(name='Calibri', size=11)

ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

# Style Header Row
for col_num in range(1, len(df.columns) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = cell_border

# Style Data Rows with Zebra Striping
for row_num in range(2, len(df) + 2):
    row_fill = zebra_fill if (row_num % 2 == 0) else white_fill
    for col_num in range(1, len(df.columns) + 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.fill = row_fill
        cell.font = cell_font
        cell.border = cell_border
        if col_num in [1, 5, 6]:  # ID, State ID, Priority ID
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col_num == 2:  # Ticket Number
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Auto-fit Column Widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 16)

# --- FULLY AUTOMATED DYNAMIC VIEW BOUNDING ---
max_col = len(df.columns)  # Automatically detects number of columns
max_row = len(df) + 1      # Automatically detects number of ticket rows

buffer_cols = 2   # Leaves 2 blank columns of breathing room
buffer_rows = 15  # Leaves 15 blank rows below for future additions

# Hide excess columns past data + buffer (up to column 100)
for col_idx in range(max_col + buffer_cols + 1, 101):
    ws.column_dimensions[get_column_letter(col_idx)].hidden = True

# Hide excess rows past data + buffer (up to row 500)
for r in range(max_row + buffer_rows + 1, 501):
    ws.row_dimensions[r].hidden = True
# ---------------------------------------------

# Watermark / Faded Logo Insertion
logo_path = '/home/ed/dhc_logo.png'
watermark_path = '/home/ed/dhc_logo_watermark.png'

if os.path.exists(logo_path):
    try:
        img_pil = Image.open(logo_path).convert('RGBA')
        alpha = img_pil.split()[3]
        alpha = alpha.point(lambda p: int(p * 0.20))  # 20% opacity watermark
        img_pil.putalpha(alpha)
        img_pil.save(watermark_path)

        img = openpyxl.drawing.image.Image(watermark_path)
        img.width = 320
        img.height = 320
        ws.add_image(img, 'D6')
    except Exception as e:
        print('Logo processing error:', e)

wb.save(excel_path)
print("Ticket export and formatting completed successfully!")
